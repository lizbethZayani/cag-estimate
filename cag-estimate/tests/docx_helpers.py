"""Build tiny valid .docx files for tests using python-docx itself."""

import io

from docx import Document


def build_docx(paragraphs: list[str] | None = None, table: list[list[str]] | None = None) -> bytes:
    """A .docx with the given paragraphs, then an optional table (rows of cell texts)."""
    document = Document()
    for paragraph in paragraphs or []:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, cell_text in enumerate(row):
                grid.cell(row_index, column_index).text = cell_text
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
